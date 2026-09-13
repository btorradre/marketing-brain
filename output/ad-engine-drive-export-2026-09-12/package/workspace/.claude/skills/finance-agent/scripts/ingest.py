#!/usr/bin/env python3
"""
Ingest — every raw transaction source into ONE table: data/transactions.json.

Sources handled (auto-detected per file):
  1. GPT 5.6 Sol raw export  (finances/inbox/*.csv|json|xlsx)  — the fetch contract in
     references/gpt-sol-fetch-prompt.md: date, account, description, amount [, merchant, currency,
     txn_id, hint, memo]. NO categories expected — categorize.py does that.
  2. Plaid-style workbook (sheet "All Expenses": account_name, vendor, name, merchant_name, amount,
     business_category, personal_finance_category_*, status) — the older ChatGPT export shape.
  3. Legacy interop CSV/JSON (date, description, amount, category, brand, source) — the `category`
     becomes a high-confidence hint, not a lock.
  4. Apple Card statement PDFs (finances/statements/*.pdf) via `pdftotext -layout`.
  5. --from-ledger: migrate entries from the old shopify-financials ledger that have NO raw source
     we can re-ingest (e.g. Meta API pulls) so nothing is lost when the ledger is rebuilt.

Every transaction: id (sha1 of account|provider txn_id when the source gives one, else date|amount|desc|account —
same-day same-amount API auto-reloads are real distinct charges), date, month, account, description, merchant,
amount (+ = money OUT), currency, source (account tag; feeds the P&L's "bank month complete" test),
source_file, hints{}, category (None until categorized), brand, cat_source, confidence, locked.

Usage:
  python3 ingest.py                       # inbox + new Apple Card PDFs (+ legacy ledger on first run)
  python3 ingest.py --inbox-only
  python3 ingest.py --file ~/Downloads/x.csv
  python3 ingest.py --apple-card          # (re)parse every Apple Card PDF
  python3 ingest.py --from-ledger         # migrate non-reingestable legacy ledger rows
  python3 ingest.py --status
"""
import os, re, sys, csv, io, json, shutil, argparse, subprocess, hashlib, datetime as dt
from collections import Counter, defaultdict
from fin_common import *  # noqa

EXCLUDED_STATUS = re.compile(r"excluded|credit card payment|internal transfer", re.I)


def load_txns():
    return load_json(TXNS, {"transactions": [], "ingested_files": {}})


# ------------------------------------------------------------------ readers
def _rows_from_csv(text):
    return list(csv.DictReader(io.StringIO(text)))


def _rows_from_json(text):
    data = json.loads(text)
    if isinstance(data, dict):
        for k in ("transactions", "entries", "expenses", "rows", "data"):
            if isinstance(data.get(k), list):
                data = data[k]; break
        else:
            data = [data]
    return data


def _rows_from_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    pick = None
    for name in ("All Expenses", "Transactions", "transactions", "Sheet1"):
        if name in wb.sheetnames:
            pick = wb[name]; break
    sheets = [pick] if pick is not None else wb.worksheets
    for ws in sheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        hdr_i = next((i for i, r in enumerate(rows[:10]) if r and sum(1 for c in r if isinstance(c, str)) >= 3
                      and any(str(c).strip().lower() in ("date", "amount", "description") for c in r if c)), None)
        if hdr_i is None: continue
        hdr = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_i]]
        for r in rows[hdr_i + 1:]:
            if r is None or all(c is None for c in r): continue
            out.append({hdr[i]: r[i] for i in range(min(len(hdr), len(r))) if hdr[i]})
        if out: break
    return out


def read_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _rows_from_xlsx(path)
    text = open(path, encoding="utf-8-sig").read()
    if ext == ".json" or text.lstrip()[:1] in "[{":
        return _rows_from_json(text)
    return _rows_from_csv(text)


# ------------------------------------------------------------------ normalizer
def normalize_row(row, source_file):
    """Any of the tabular shapes -> canonical transaction dict (or None + reason)."""
    r = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
    date = parse_date(r.get("date") or r.get("transaction_date") or r.get("posted") or r.get("txn_date"))
    if not date:
        return None, "no date"
    amount = parse_amount(r.get("amount"))
    if amount is None:
        return None, "no amount"
    direction = str(r.get("direction") or r.get("type") or "").strip().lower()
    if direction in ("in", "credit", "deposit", "inflow", "income") and amount > 0:
        amount = -amount           # money IN is negative in our convention (positive = outflow)
    elif direction in ("out", "debit", "outflow", "expense") and amount < 0:
        amount = -amount
    account = (r.get("account") or r.get("account_name") or r.get("source") or r.get("card") or "unknown")
    desc = (r.get("description") or r.get("name") or r.get("memo") or r.get("merchant_name") or r.get("vendor") or "")
    merchant = (r.get("merchant") or r.get("vendor") or r.get("merchant_name") or "")
    desc, merchant = str(desc).strip(), str(merchant).strip()
    if not desc and not merchant:
        return None, "no description"
    if merchant and merchant.lower() != desc.lower() and merchant.lower() not in desc.lower():
        full_desc = f"{merchant} ; {desc}" if desc else merchant
    else:
        full_desc = desc or merchant
    hints = {}
    for k in ("business_category", "personal_finance_category_primary", "personal_finance_category_detailed",
              "personal_finance_category_confidence_level", "status", "hint", "plaid_category", "category_hint", "memo", "txn_id", "currency"):
        if r.get(k) not in (None, ""):
            hints[k] = str(r[k]).strip()
    # a `category` column (legacy interop shape) is a strong hint but never a lock
    if r.get("category") not in (None, ""):
        c = norm_category(r["category"])
        hints["provided_category"] = c or str(r["category"]).strip()
    if EXCLUDED_STATUS.search(hints.get("status", "")):
        hints["excluded_by_source"] = "true"
    brand = str(r.get("brand") or "").strip().title() or None
    if brand and brand not in BRANDS:
        brand = None
    tag = account_tag(account)
    provider_id = str(r.get("txn_id") or r.get("transaction_id") or "").strip()
    tid = hashlib.sha1(f"{tag}|{provider_id}".encode()).hexdigest()[:16] if provider_id else txn_id(date, amount, full_desc, tag)
    t = {"id": tid, "date": date, "month": month_of(date),
         "account": str(account), "source": tag, "description": full_desc,
         "merchant": clean_merchant(merchant or desc), "amount": round(amount, 2),
         "currency": hints.pop("currency", "USD"), "source_file": os.path.basename(source_file),
         "hints": hints, "category": None, "brand": brand, "cat_source": None, "confidence": None, "locked": False}
    return t, None


# ------------------------------------------------------------------ Apple Card PDF
APPLE_TXN = re.compile(r"^\s*(\d{2}/\d{2}/\d{4})\s+(.+?)\s{2,}(\d+%\s+\$[\d,]+\.\d{2}\s+)?(-?\$[\d,]+\.\d{2})\s*$")


def parse_apple_card(pdf_path):
    """Statement PDF -> list of transactions. Skips payments (ACH) and Daily Cash adjustments."""
    txt = subprocess.run(["pdftotext", "-layout", pdf_path, "-"], capture_output=True, text=True).stdout
    section = None
    out, pending = [], None
    for line in txt.splitlines():
        s = line.strip()
        if s.startswith("Payments") and "Date" not in s: section = "payments"; continue
        if s.startswith("Transactions"): section = "transactions"; continue
        if s.startswith("Daily Cash Adjustments") or s.startswith("Interest Charged") or s.startswith("Installments"):
            section = s.split()[0].lower(); continue
        if s.startswith("Total ") or s.startswith("Apple Card is issued"): continue
        m = APPLE_TXN.match(line)
        if m and section == "transactions":
            d = dt.datetime.strptime(m.group(1), "%m/%d/%Y").date().isoformat()
            desc = m.group(2).strip()
            amt = parse_amount(m.group(4))
            pending = {"date": d, "desc": desc, "amount": amt}
            out.append(pending); continue
        if m and section == "interest":
            d = dt.datetime.strptime(m.group(1), "%m/%d/%Y").date().isoformat()
            out.append({"date": d, "desc": "Apple Card interest charge ; " + m.group(2).strip(), "amount": parse_amount(m.group(4))})
            pending = None; continue
        # wrapped description continuation (indented, no date, no amount)
        if pending and section == "transactions" and s and not re.match(r"^\d{2}/\d{2}/\d{4}", s) and "$" not in s and len(s) < 60:
            pending["desc"] += " " + s
        else:
            pending = None
    txns = []
    for o in out:
        if o["amount"] is None: continue
        t = {"id": txn_id(o["date"], o["amount"], o["desc"], "apple_card"), "date": o["date"], "month": month_of(o["date"]),
             "account": "Apple Card", "source": "apple_card", "description": o["desc"], "merchant": clean_merchant(o["desc"]),
             "amount": round(o["amount"], 2), "currency": "USD", "source_file": os.path.basename(pdf_path),
             "hints": {}, "category": None, "brand": None, "cat_source": None, "confidence": None, "locked": False}
        txns.append(t)
    return txns


# ------------------------------------------------------------------ legacy ledger migration
REINGESTABLE = ("business_gold", "business_platinum", "checking", "business_checking", "apple_card")


def migrate_legacy_ledger():
    """Old ledger rows whose raw source we can't re-read (meta_api:*, manual, etc.) become locked transactions."""
    led = load_json(LEDGER, {"entries": []})
    keep = []
    for e in led.get("entries", []):
        src = e.get("source") or ""
        if src.startswith(REINGESTABLE):
            continue
        date = e.get("date") or (e["month"] + "-01")
        t = {"id": txn_id(date, e["amount"], e["description"], src or "legacy"), "date": date, "month": e["month"],
             "account": src or "legacy", "source": src or "legacy", "description": e["description"],
             "merchant": clean_merchant(e["description"]), "amount": e["amount"], "currency": "USD",
             "source_file": "legacy expenses.json", "hints": {}, "category": e["category"], "brand": e.get("brand"),
             "cat_source": "legacy", "confidence": 1.0, "locked": True}
        keep.append(t)
    return keep


# ------------------------------------------------------------------ main
def add_txns(store, new, label):
    existing = {t["id"] for t in store["transactions"]}
    added = [t for t in new if t["id"] not in existing]
    store["transactions"].extend(added)
    print(f"  {label}: {len(new)} parsed, {len(added)} new, {len(new) - len(added)} duplicates skipped")
    return len(added)


def ingest_file(store, path, move=True):
    rows = read_file(path)
    new, bad = [], Counter()
    for row in rows:
        t, why = normalize_row(row, path)
        if t: new.append(t)
        else: bad[why] += 1
    n = add_txns(store, new, os.path.basename(path))
    if bad: print(f"    skipped rows: {dict(bad)}")
    store["ingested_files"][os.path.basename(path)] = {"at": now_iso(), "rows": len(rows), "added": n}
    if move and os.path.dirname(os.path.abspath(path)) == INBOX:
        os.makedirs(PROCESSED, exist_ok=True)
        shutil.move(path, os.path.join(PROCESSED, os.path.basename(path)))
    return n


def print_status(store):
    T = store["transactions"]
    print(f"transactions.json: {len(T)} transactions")
    by_src = defaultdict(lambda: [0, set()])
    for t in T:
        by_src[t["source"]][0] += 1; by_src[t["source"]][1].add(t["month"])
    for s, (n, ms) in sorted(by_src.items()):
        print(f"  {s:28} {n:5d}   {min(ms)} .. {max(ms)}")
    unc = sum(1 for t in T if not t["category"])
    print(f"uncategorized: {unc}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--file", action="append", help="ingest a specific file (not moved)")
    ap.add_argument("--inbox-only", action="store_true")
    ap.add_argument("--apple-card", action="store_true", help="(re)parse every Apple Card PDF in finances/statements")
    ap.add_argument("--from-ledger", action="store_true", help="migrate legacy ledger rows with no raw source")
    ap.add_argument("--status", action="store_true")
    ap.add_argument("--no-move", action="store_true", help="leave inbox files in place")
    args = ap.parse_args()

    store = load_txns()
    if args.status:
        print_status(store); return

    total = 0
    if args.file:
        for f in args.file:
            total += ingest_file(store, os.path.expanduser(f), move=False)
    else:
        files = sorted(f for f in os.listdir(INBOX) if os.path.isfile(os.path.join(INBOX, f))
                       and f.lower().endswith((".csv", ".json", ".xlsx", ".xlsm")) and not f.startswith("."))
        if files:
            print(f"Inbox: {len(files)} file(s)")
            for f in files:
                total += ingest_file(store, os.path.join(INBOX, f), move=not args.no_move)
        else:
            print("Inbox: empty")
        if not args.inbox_only:
            pdfs = sorted(f for f in os.listdir(STATEMENTS) if f.lower().endswith(".pdf")) if os.path.isdir(STATEMENTS) else []
            new_pdfs = [f for f in pdfs if args.apple_card or f not in store["ingested_files"]]
            if new_pdfs:
                print(f"Apple Card statements: {len(new_pdfs)} to parse")
                for f in new_pdfs:
                    tx = parse_apple_card(os.path.join(STATEMENTS, f))
                    n = add_txns(store, tx, f); total += n
                    store["ingested_files"][f] = {"at": now_iso(), "rows": len(tx), "added": n}
            if args.from_ledger or "legacy expenses.json" not in store["ingested_files"]:
                legacy = migrate_legacy_ledger()
                if legacy:
                    n = add_txns(store, legacy, "legacy ledger (non-reingestable rows)"); total += n
                store["ingested_files"]["legacy expenses.json"] = {"at": now_iso(), "rows": len(legacy), "added": len(legacy)}

    store["transactions"].sort(key=lambda t: (t["date"], t["source"], t["description"]))
    store["updated_at"] = now_iso()
    save_json(TXNS, store)
    print(f"\n+{total} transactions → {TXNS}")
    print_status(store)


if __name__ == "__main__":
    main()
